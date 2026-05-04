"""
pdf_engine.py — Thrust Test Rig PDF Report Generator
Builds a styled reportlab PDF from a data dict + chart images.
Edit this file to change PDF layout, colours, sections, or formatting.
"""

import io
import re
import matplotlib.pyplot as plt
from pathlib import Path

from reportlab.lib.pagesizes import A4
from reportlab.lib import colors as rl_colors
from reportlab.lib.units import mm
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_LEFT, TA_CENTER
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
    Image, PageBreak, KeepTogether
)

# ── Paths (resolved relative to this file's location) ──
BASE_DIR      = Path(__file__).parent
TEMPLATE_PATH = BASE_DIR / "Format.xlsx"
LOGO_PATH     = BASE_DIR / "ideaforge-logo.jpeg"

# ── Colours ──
NAVY  = rl_colors.HexColor("#1B5E20")   # IdeaForge dark green
BLUE  = rl_colors.HexColor("#2E7D32")   # IdeaForge mid green
MGRAY = rl_colors.HexColor("#CCCCCC")
LGRAY = rl_colors.HexColor("#F5F5F5")
WHITE = rl_colors.white
BLACK = rl_colors.black


def _S(nm, **kw):
    return ParagraphStyle(nm, **kw)


def _banner(txt, sty, bg, W, pad=7):
    t = Table([[Paragraph(txt, sty)]], colWidths=[W])
    t.setStyle(TableStyle([
        ("BACKGROUND",    (0,0),(-1,-1), bg),
        ("TOPPADDING",    (0,0),(-1,-1), pad),
        ("BOTTOMPADDING", (0,0),(-1,-1), pad),
        ("LEFTPADDING",   (0,0),(-1,-1), 10),
    ]))
    return t


def _resolve(data, text):
    """Replace all {{key}} placeholders in text with values from data dict."""
    def _repl(m):
        key = m.group(1).strip()
        val = data.get(key, "")
        try:
            fv = float(val)
            if key in ("max_rpm", "target_rpm", "num_rows"):
                return f"{int(fv):,}"
            elif key in ("duration_s", "time_at_target_rpm"):
                return f"{fv:.1f}"
            elif key == "overall_efficiency":
                return f"{fv:.4f}"
            elif key == "mechanical_efficiency":
                return f"{fv:.2f}"
            else:
                return f"{fv:.3f}".rstrip("0").rstrip(".")
        except (ValueError, TypeError):
            return str(val) if val not in (None, "") else "—"
    return re.sub(r"\{\{(\w+)\}\}", _repl, str(text))


def build_pdf_report(data, chart_images, run_name, template_path=None):
    """
    Build and return PDF bytes.

    Args:
        data:          dict of placeholder_key → value
        chart_images:  list of (title_str, png_bytes)
        run_name:      string shown in the header banner
        template_path: optional Path override for Format.xlsx

    Returns:
        bytes — the complete PDF
    """
    tpl = template_path or TEMPLATE_PATH

    # ── Fill Format.xlsx placeholders (for reference / future LibreOffice use) ──
    if tpl.exists():
        import openpyxl as _oxl
        wb = _oxl.load_workbook(tpl)
        ws = wb.active
        for row in ws.iter_rows():
            for cell in row:
                if cell.value and isinstance(cell.value, str) and "{{" in cell.value:
                    cell.value = _resolve(data, cell.value)

    # ── Build PDF ──
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4,
        leftMargin=12*mm, rightMargin=12*mm,
        topMargin=12*mm, bottomMargin=12*mm,
        title=run_name)
    W = A4[0] - 24*mm

    # Styles
    sty_h1   = _S("h1",   fontName="Helvetica-Bold",   fontSize=14, textColor=WHITE,  alignment=TA_CENTER)
    sty_meta = _S("meta", fontName="Helvetica",         fontSize=8,  textColor=rl_colors.HexColor("#AAAAAA"), alignment=TA_CENTER)
    sty_sec  = _S("sec",  fontName="Helvetica-Bold",    fontSize=9,  textColor=WHITE,  alignment=TA_LEFT)
    sty_lbl  = _S("lbl",  fontName="Helvetica-Bold",    fontSize=8,  textColor=rl_colors.HexColor("#222222"))
    sty_grp  = _S("grp",  fontName="Helvetica-Oblique", fontSize=7,  textColor=rl_colors.HexColor("#666666"))
    sty_val  = _S("val",  fontName="Helvetica",         fontSize=8,  textColor=rl_colors.HexColor("#1A73E8"))
    sty_obs  = _S("obs",  fontName="Helvetica",         fontSize=8,  textColor=BLACK,  leading=13)
    sty_chk  = _S("chk",  fontName="Helvetica",         fontSize=7,  textColor=rl_colors.HexColor("#333333"))
    sty_res  = _S("res",  fontName="Helvetica-Bold",    fontSize=8,  textColor=NAVY)
    sty_cht  = _S("cht",  fontName="Helvetica-Bold",    fontSize=9,  textColor=NAVY,   spaceAfter=3)
    sty_sub2 = _S("sub2", fontName="Helvetica",         fontSize=8,  textColor=rl_colors.HexColor("#A5D6A7"), alignment=TA_LEFT)

    def pv(key):
        """Paragraph with formatted value from data dict."""
        v = data.get(key, "—")
        try:
            fv = float(v)
            if key in ("max_rpm", "target_rpm"):       v = f"{int(fv):,}"
            elif key == "overall_efficiency":           v = f"{fv:.4f}"
            elif key == "mechanical_efficiency":        v = f"{fv:.2f}"
            elif key in ("duration_s", "time_at_target_rpm"): v = f"{fv:.1f}"
            else:                                       v = f"{fv:.3f}".rstrip("0").rstrip(".")
        except Exception:
            pass
        return Paragraph(str(v) if v not in (None, "") else "—", sty_val)

    story = []

    # ── HEADER ──
    if LOGO_PATH.exists():
        logo_cell = Image(str(LOGO_PATH), width=16*mm, height=16*mm)
    else:
        logo_cell = Paragraph("", sty_h1)

    hdr = Table(
        [[logo_cell,
          Paragraph(run_name, sty_h1),
          Paragraph(f"Thrust Test Report &nbsp;·&nbsp; {data.get('test_date','')}", sty_sub2)]],
        colWidths=[20*mm, W * 0.62, W - 20*mm - W * 0.62]
    )
    hdr.setStyle(TableStyle([
        ("BACKGROUND",    (0,0), (-1,-1), NAVY),
        ("BACKGROUND",    (0,0), (0,0),   WHITE),
        ("VALIGN",        (0,0), (-1,-1), "MIDDLE"),
        ("LEFTPADDING",   (0,0), (0,0),   3),
        ("RIGHTPADDING",  (0,0), (0,0),   3),
        ("LEFTPADDING",   (1,0), (1,0),   10),
        ("LEFTPADDING",   (2,0), (2,0),   6),
        ("TOPPADDING",    (0,0), (-1,-1), 6),
        ("BOTTOMPADDING", (0,0), (-1,-1), 6),
    ]))
    story.append(hdr)

    # ── META STRIP ──
    meta = Table([[
        Paragraph(f"<b>File:</b> {data.get('filename','')}", sty_meta),
        Paragraph(f"<b>Date:</b> {data.get('test_date','')}", sty_meta),
        Paragraph(f"<b>Duration:</b> {data.get('duration_s','')}s  |  <b>Rows:</b> {data.get('num_rows','')}", sty_meta),
    ]], colWidths=[W*0.45, W*0.20, W*0.35])
    meta.setStyle(TableStyle([
        ("BACKGROUND",    (0,0),(-1,-1), LGRAY),
        ("TOPPADDING",    (0,0),(-1,-1), 4),
        ("BOTTOMPADDING", (0,0),(-1,-1), 4),
        ("LEFTPADDING",   (0,0),(-1,-1), 6),
        ("LINEBELOW",     (0,0),(-1,-1), 0.5, MGRAY),
    ]))
    story.append(meta)
    story.append(Spacer(1, 4*mm))

    # ── INITIAL PARAMETERS ──
    CW = [W*0.18, W*0.32, W*0.38, W*0.12]
    ip_rows = [
        ("Reservoir",   "Capacity",              "res_capacity",         "L"),
        ("",            "Composition",           "res_composition",      ""),
        ("",            "Temperature",           "res_temperature",      "°C"),
        ("Duty Cycle",  "Duty Cycle & Flowrate", "duty_cycle",           ""),
        ("Temperature", "Initial ESC Temp",      "init_esc_temp",        "°C"),
        ("",            "Initial Motor Temp",    "init_motor_temp",      "°C"),
        ("",            "Ambient",               "ambient_temp",         "°C"),
        ("",            "ESC Inlet Coolant",     "esc_inlet_coolant",    "°C"),
        ("",            "Motor Inlet Coolant",   "motor_inlet_coolant",  "°C"),
        ("Flowrate",    "ESC Inlet",             "esc_inlet_flow",       "LPM"),
        ("",            "Motor Inlet",           "motor_inlet_flow",     "LPM"),
        ("Pressure",    "ESC Inlet",             "esc_inlet_pressure",   "Bar"),
        ("",            "Motor Inlet",           "motor_inlet_pressure", "Bar"),
        ("Battery",     "Battery Voltage",       "battery_voltage",      "V"),
        ("",            "SOC",                   "battery_soc",          ""),
        ("",            "SOH",                   "battery_soh",          ""),
        ("Fintube",     "Inlet Temperature",     "fin_inlet_temp",       "°C"),
        ("",            "Outlet Temperature",    "fin_outlet_temp",      "°C"),
        ("Target RPM",  "",                      "target_rpm",           "RPM"),
    ]
    ip_data = []
    ip_cmds = [
        ("GRID",          (0,0),(-1,-1), 0.4, MGRAY),
        ("VALIGN",        (0,0),(-1,-1), "MIDDLE"),
        ("TOPPADDING",    (0,0),(-1,-1), 3),
        ("BOTTOMPADDING", (0,0),(-1,-1), 3),
        ("LEFTPADDING",   (0,0),(-1,-1), 5),
    ]
    for grp, lbl, key, unit in ip_rows:
        ip_data.append([
            Paragraph(grp,  sty_grp),
            Paragraph(lbl,  sty_lbl),
            pv(key),
            Paragraph(unit, sty_chk),
        ])
    t_ip = Table(ip_data, colWidths=CW)
    t_ip.setStyle(TableStyle(ip_cmds))
    story.append(_banner("INITIAL PARAMETERS", sty_sec, BLUE, W, pad=5))
    story.append(t_ip)
    story.append(Spacer(1, 4*mm))

    # ── RESULTS ──
    res_rows = [
        ("Max. Temp — ESC Inlet",      "max_esc_inlet_temp",    "°C"),
        ("Max. Temp — Motor Inlet",    "max_motor_inlet_temp",  "°C"),
        ("Max. Pressure — ESC Inlet",  "max_esc_pressure",      "Bar"),
        ("Battery Voltage (post-run)", "battery_voltage_post",  "V"),
        ("Max. RPM",                   "max_rpm",               "RPM"),
        ("Max. Torque",                "max_torque",            "Nm"),
        ("Max. Thrust",                "max_thrust",            "N"),
        ("Fin Tube Inlet Temp (max)",  "max_fin_inlet_temp",    "°C"),
        ("Fin Tube Outlet Temp (max)", "max_fin_outlet_temp",   "°C"),
        ("Max. ESC Temp",              "max_esc_temp",          "°C"),
        ("Max. Motor Temp",            "max_motor_temp",        "°C"),
        ("Time at Target RPM",         "time_at_target_rpm",    "s"),
        ("Mechanical Power",           "mechanical_power",      "W"),
        ("Electrical Power",           "electrical_power",      "W"),
        ("Mechanical Efficiency",      "mechanical_efficiency", "%"),
        ("Overall Efficiency",         "overall_efficiency",    "g/W"),
    ]
    CW2 = [W*0.55, W*0.33, W*0.12]
    res_data = []
    res_cmds = [
        ("GRID",          (0,0),(-1,-1), 0.4, MGRAY),
        ("VALIGN",        (0,0),(-1,-1), "MIDDLE"),
        ("TOPPADDING",    (0,0),(-1,-1), 3),
        ("BOTTOMPADDING", (0,0),(-1,-1), 3),
        ("LEFTPADDING",   (0,0),(-1,-1), 5),
        ("BACKGROUND",    (0,0),(0,-1),  LGRAY),
    ]
    for lbl, key, unit in res_rows:
        res_data.append([
            Paragraph(lbl,  sty_res),
            pv(key),
            Paragraph(unit, sty_chk),
        ])
    t_res = Table(res_data, colWidths=CW2)
    t_res.setStyle(TableStyle(res_cmds))
    story.append(KeepTogether([_banner("RESULTS", sty_sec, NAVY, W, pad=5), t_res]))
    story.append(Spacer(1, 4*mm))

    # ── OBSERVATIONS ──
    obs = data.get("notes", "")
    if obs:
        obs_tbl = Table([[Paragraph(str(obs), sty_obs)]], colWidths=[W])
        obs_tbl.setStyle(TableStyle([
            ("GRID",          (0,0),(-1,-1), 0.4, MGRAY),
            ("TOPPADDING",    (0,0),(-1,-1), 6),
            ("BOTTOMPADDING", (0,0),(-1,-1), 6),
            ("LEFTPADDING",   (0,0),(-1,-1), 8),
        ]))
        story.append(_banner("OBSERVATIONS", sty_sec, BLUE, W, pad=5))
        story.append(obs_tbl)
        story.append(Spacer(1, 4*mm))

    # ── CHARTS ──
    if chart_images:
        story.append(PageBreak())
        story.append(_banner("TEST CHARTS", sty_sec, NAVY, W))
        story.append(Spacer(1, 4*mm))
        for chart_title, png_bytes in chart_images:
            story.append(Paragraph(chart_title, sty_cht))
            img = Image(io.BytesIO(png_bytes), width=W, height=W * 0.38)
            img.hAlign = "LEFT"
            story.append(img)
            story.append(Spacer(1, 4*mm))

    doc.build(story)
    buf.seek(0)
    return buf.read()


def fig_to_png(fig, dpi=150):
    """Convert a matplotlib figure to PNG bytes and close it."""
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=dpi, bbox_inches="tight",
                facecolor=fig.get_facecolor())
    plt.close(fig)
    buf.seek(0)
    return buf.read()
